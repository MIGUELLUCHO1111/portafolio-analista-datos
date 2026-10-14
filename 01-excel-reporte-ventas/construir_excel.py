"""Genera reporte_ventas.xlsx con tablas, fórmulas reales, formato condicional y gráfico.
Las fórmulas se guardan en inglés (formato interno de Excel); Excel en español las muestra
automáticamente como BUSCARV, SUMAR.SI.CONJUNTO, SI, etc. Uso: python construir_excel.py"""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import date

AQUI = Path(__file__).parent
DATOS = AQUI.parent / "datos"
leer = lambda n: list(csv.DictReader(open(DATOS / f"{n}.csv", encoding="utf-8")))
ventas, productos, vendedores = leer("ventas"), leer("productos"), leer("vendedores")
P = {p["id"]: p for p in productos}
V = {v["id"]: v for v in vendedores}

wb = Workbook()
TITULO = Font(bold=True, size=14, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="0A1030")
ENC = Font(bold=True)

def tabla(ws, nombre, ref, estilo="TableStyleMedium2"):
    t = Table(displayName=nombre, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=estilo, showRowStripes=True)
    ws.add_table(t)

# ---------- Productos ----------
wp = wb.active; wp.title = "Productos"
wp.append(["Producto", "Categoría", "Precio", "Costo"])
for p in productos:
    wp.append([p["nombre"], p["categoria"], float(p["precio"]), float(p["costo"])])
tabla(wp, "tProductos", f"A1:D{len(productos) + 1}")

# ---------- Vendedores ----------
wv = wb.create_sheet("Vendedores")
wv.append(["Vendedor", "Región", "Meta"])
for v in vendedores:
    wv.append([v["nombre"], v["region"], float(v["meta"])])
tabla(wv, "tVendedores", f"A1:C{len(vendedores) + 1}")

# ---------- Datos (tabla de hechos con columnas calculadas) ----------
wd = wb.create_sheet("Datos", 0)
cols = ["Fecha", "Región", "Canal", "Vendedor", "Producto", "Cantidad", "Precio", "Categoría", "Costo", "Total", "Margen", "Mes", "Tamaño"]
wd.append(cols)
for i, s in enumerate(ventas, start=2):
    y, m, d = map(int, s["fecha"].split("-"))
    wd.append([date(y, m, d), s["region"], s["canal"], V[s["vendedor_id"]]["nombre"], P[s["producto_id"]]["nombre"],
               int(s["cantidad"]), float(s["precio_unit"]),
               f"=VLOOKUP(E{i},Productos!$A$2:$D$11,2,FALSE)",            # BUSCARV
               f"=VLOOKUP(E{i},Productos!$A$2:$D$11,4,FALSE)",
               f"=F{i}*G{i}",
               f"=F{i}*(G{i}-I{i})",
               f"=MONTH(A{i})",                                  # MES
               f'=IF(J{i}>=1000,"Grande",IF(J{i}>=300,"Mediana","Pequeña"))'])  # SI anidado
    wd[f"A{i}"].number_format = "yyyy-mm-dd"
n = len(ventas) + 1
tabla(wd, "tVentas", f"A1:M{n}")
for c, w in zip("ABCDEFGHIJKLM", [12, 10, 9, 15, 20, 10, 10, 13, 9, 11, 11, 6, 10]):
    wd.column_dimensions[c].width = w
wd.freeze_panes = "A2"
wd.conditional_formatting.add(f"J2:J{n}", DataBarRule(start_type="min", end_type="max", color="1E5BFF"))
wd.conditional_formatting.add(f"J2:J{n}", CellIsRule(operator="greaterThanOrEqual", formula=["1000"], font=Font(bold=True, color="11A64B")))

# ---------- Resumen (KPIs y análisis con fórmulas) ----------
wr = wb.create_sheet("Resumen", 0)
wr["A1"] = "Reporte de ventas · enero–junio 2026"; wr["A1"].font = TITULO
for c in "ABCDEFG":
    wr[f"{c}1"].fill = FONDO
wr["A2"] = "Todas las cifras se calculan con fórmulas sobre la hoja Datos."
kpis = [
    ("Ingreso total", f"=SUM(Datos!J2:J{n})"),                                   # SUMA
    ("Margen total", f"=SUM(Datos!K2:K{n})"),
    ("Margen %", "=B5/B4"),
    ("N° de ventas", f"=COUNTA(Datos!A2:A{n})"),                                 # CONTARA
    ("Ticket promedio", f"=AVERAGE(Datos!J2:J{n})"),                             # PROMEDIO
    ("Venta máxima", f"=MAX(Datos!J2:J{n})"),
    ("Ventas grandes (≥ 1000)", f'=COUNTIF(Datos!J2:J{n},">=1000")'),           # CONTAR.SI
]
wr["A3"], wr["B3"] = "Indicador", "Valor"
for i, (k, f) in enumerate(kpis, start=4):
    wr[f"A{i}"], wr[f"B{i}"] = k, f
wr["B6"].number_format = "0.0%"
for r in (4, 5, 8, 9):
    wr[f"B{r}"].number_format = "#,##0.00"

wr["D3"], wr["E3"], wr["F3"], wr["G3"] = "Región", "Ingreso", "Participación", "Ventas Online"
for i, reg in enumerate(["Norte", "Sur", "Este", "Oeste"], start=4):
    wr[f"D{i}"] = reg
    wr[f"E{i}"] = f'=SUMIF(Datos!B2:B{n},D{i},Datos!J2:J{n})'                      # SUMAR.SI
    wr[f"F{i}"] = f"=E{i}/$B$4"                                                     # referencia absoluta
    wr[f"G{i}"] = f'=COUNTIFS(Datos!B2:B{n},D{i},Datos!C2:C{n},"Online")'          # CONTAR.SI.CONJUNTO
    wr[f"E{i}"].number_format = "#,##0"; wr[f"F{i}"].number_format = "0.0%"

wr["A13"], wr["B13"], wr["C13"], wr["D13"], wr["E13"] = "Vendedor", "Ingreso", "Meta", "Cumplimiento", "Estado"
for i, v in enumerate(vendedores, start=14):
    wr[f"A{i}"] = v["nombre"]
    wr[f"B{i}"] = f'=SUMIFS(Datos!J2:J{n},Datos!D2:D{n},A{i})'                       # SUMAR.SI.CONJUNTO
    wr[f"C{i}"] = f"=VLOOKUP(A{i},Vendedores!$A$2:$C$7,3,FALSE)"
    wr[f"D{i}"] = f"=IFERROR(B{i}/C{i},0)"                                          # SI.ERROR
    wr[f"E{i}"] = f'=IF(D{i}>=1,"Cumplió","Por debajo")'
    wr[f"B{i}"].number_format = wr[f"C{i}"].number_format = "#,##0"; wr[f"D{i}"].number_format = "0.0%"
last = 13 + len(vendedores)
wr.conditional_formatting.add(f"D14:D{last}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor="C6EFCE")))
wr.conditional_formatting.add(f"D14:D{last}", CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor="FFC7CE")))
for cell in ("A3", "B3", "D3", "E3", "F3", "G3", "A13", "B13", "C13", "D13", "E13"):
    wr[cell].font = ENC
for c, w in zip("ABCDEFG", [26, 14, 10, 14, 12, 13, 13]):
    wr.column_dimensions[c].width = w

graf = BarChart(); graf.title = "Ingreso por región"; graf.y_axis.title = "USD"; graf.height, graf.width = 7.5, 14
graf.add_data(Reference(wr, min_col=5, min_row=3, max_row=7), titles_from_data=True)
graf.set_categories(Reference(wr, min_col=4, min_row=4, max_row=7))
wr.add_chart(graf, "D13" if False else "G13")

# ---------- Tu turno ----------
wt = wb.create_sheet("Tu turno")
tareas = [
    "1. En Datos, convierte la columna Mes en nombre del mes con =TEXTO(A2;\"mmmm\").",
    "2. Crea una tabla dinámica (Insertar → Tabla dinámica) con Producto en filas, Canal en columnas y Total en valores.",
    "3. Agrega una segmentación por Región a esa tabla dinámica.",
    "4. En Resumen, calcula el ingreso de 'Laptop Pro 14' en el Norte con SUMAR.SI.CONJUNTO.",
    "5. Usa BUSCARX para traer la región de cada vendedor en la hoja Datos.",
    "6. Crea un gráfico de líneas del ingreso por mes.",
    "7. Guarda, haz commit y push a GitHub: git add . && git commit -m \"Excel: tabla dinámica y gráficos\" && git push",
]
wt["A1"] = "Tareas para practicar (aprender haciendo)"; wt["A1"].font = Font(bold=True, size=13)
for i, t in enumerate(tareas, start=3):
    wt[f"A{i}"] = t; wt[f"A{i}"].alignment = Alignment(wrap_text=True)
wt.column_dimensions["A"].width = 110

wb.calculation.fullCalcOnLoad = True
wb.save(AQUI / "reporte_ventas.xlsx")
print("reporte_ventas.xlsx creado")
